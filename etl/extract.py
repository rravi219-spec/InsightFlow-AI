"""Read-only streaming extraction; source workbook is never modified."""
from datetime import date, datetime
from hashlib import sha256
import json

from openpyxl import load_workbook

from .retail_config import COLUMNS, SHEETS


def file_digest(path):
    digest = sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def raw_json(values):
    # Excel date cells retain their type and full timestamp in the lineage payload.
    return json.dumps([
        {"type": type(v).__name__, "value": v.isoformat() if isinstance(v, (date, datetime)) else v}
        for v in values
    ], ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def iter_source(path, sheets=SHEETS):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if set(book.sheetnames) != set(sheets):
            raise ValueError(f"Expected sheets {sheets}; found {book.sheetnames}")
        for sheet in sheets:
            rows = book[sheet].iter_rows(values_only=True)
            header = tuple(next(rows, ()))
            if header != COLUMNS:
                raise ValueError(f"Invalid schema in {sheet}: {header}; expected {COLUMNS}")
            for source_row, values in enumerate(rows, start=2):
                # Even entirely blank rows are accounted for, never silently skipped.
                yield sheet, source_row, dict(zip(COLUMNS, values))
    finally:
        book.close()
