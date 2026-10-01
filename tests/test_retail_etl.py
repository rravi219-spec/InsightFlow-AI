from contextlib import closing
from datetime import datetime
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest

from openpyxl import Workbook

from etl.extract import iter_source, raw_json
from etl.retail_config import COLUMNS, SHEETS
from etl.run_retail_etl import run_etl
from etl.reconcile import reconcile
from etl.transform import transform_record
from etl.validate import classify, parse_date


class RetailETLTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.xlsx"
        self.db = self.root / "retail.sqlite3"
        self.report_path = self.root / "report.json"
        self.base = [123456, "P1", "Product one", 2, datetime(2010, 1, 2, 3, 4), 1.235, 12345, "UK"]

    def write(self, rows=None, header=COLUMNS, sheets=SHEETS):
        book = Workbook()
        book.remove(book.active)
        for i, name in enumerate(sheets):
            sheet = book.create_sheet(name)
            sheet.append(header)
            for row in (rows if rows is not None else [self.base]) if i == 0 else []:
                sheet.append(row)
        book.save(self.source)
        book.close()

    def run_pipeline(self):
        return run_etl(self.source, self.db, self.report_path)

    def normalize(self, **changes):
        row = dict(zip(COLUMNS, self.base))
        row.update(changes)
        return transform_record(SHEETS[0], 2, row, raw_json(list(row.values())), "test", False, classify(row))

    def test_schema_validation(self):
        self.write(header=("Wrong",) + COLUMNS[1:])
        with self.assertRaisesRegex(ValueError, "Invalid schema"):
            list(iter_source(self.source))

    def test_both_periods_required(self):
        self.write(sheets=SHEETS[:1])
        with self.assertRaisesRegex(ValueError, "Expected sheets"):
            list(iter_source(self.source))

    def test_rows_from_both_periods_have_distinct_lineage(self):
        book = Workbook()
        book.remove(book.active)
        for name in SHEETS:
            sheet = book.create_sheet(name)
            sheet.append(COLUMNS)
            sheet.append(self.base)
        book.save(self.source)
        book.close()
        report = self.run_pipeline()
        self.assertEqual(report["sheets"], dict.fromkeys(SHEETS, 1))
        self.assertEqual(report["duplicate_source_rows"], 1)
        self.assertEqual(report["valid_analytical_rows"], 2)
        with closing(sqlite3.connect(self.db)) as conn:
            self.assertEqual(conn.execute("SELECT source_sheet,source_row FROM fact_transaction ORDER BY source_sheet").fetchall(),
                             [(name, 2) for name in SHEETS])

    def test_date_parsing(self):
        self.assertEqual(parse_date("2010-01-02T03:04:00"), self.base[4])
        for value in ["02/01/2010", "bad", 40200, None, "2010-01-01T00:00:00+01:00"]:
            self.assertIsNone(parse_date(value))
        record, reasons = self.normalize(InvoiceDate="bad")
        self.assertIsNone(record[10])
        self.assertIn("quality:invalid_dates", reasons)

    def test_exact_amount(self):
        record, reasons = self.normalize()
        self.assertFalse(reasons)
        self.assertEqual(record[12:14], (1235000, 2470000))
        record, reasons = self.normalize(Quantity=-3, Price=0.001)
        self.assertEqual(record[13], -3000)
        self.assertFalse(reasons)

    def test_cancellation_retains_source_sign(self):
        record, reasons = self.normalize(Invoice="c123456", Quantity=-2)
        self.assertEqual(record[5], "C123456")
        self.assertEqual(record[14], 1)
        self.assertEqual(record[13], -2470000)
        self.assertFalse(reasons)

    def test_missing_customer_retained(self):
        record, reasons = self.normalize(**{"Customer ID": None})
        self.assertIsNone(record[8])
        self.assertFalse(reasons)
        self.assertIn("missing_customer_ids", json.loads(record[15]))

    def test_invalid_quantities(self):
        for value in [None, "bad", 1.5, "inf", 2**40]:
            with self.subTest(value=value):
                _, reasons = self.normalize(Quantity=value)
                self.assertIn("quality:invalid_quantities", reasons)
        _, reasons = self.normalize(Quantity=0)
        self.assertIn("business:zero_quantities", reasons)

    def test_invalid_prices(self):
        for value in [None, "bad", "inf"]:
            _, reasons = self.normalize(Price=value)
            self.assertIn("quality:invalid_prices", reasons)
        for value, reason in [(0, "business:zero_prices"), (-1, "business:negative_prices"),
                              (0.0000001, "quality:unsupported_price_precision_or_range")]:
            _, reasons = self.normalize(Price=value)
            self.assertIn(reason, reasons)

    def test_nonstandard_invoice_and_missing_product(self):
        _, reasons = self.normalize(Invoice="A123456", StockCode=None)
        self.assertIn("quality:invalid_invoice_ids", reasons)
        self.assertIn("quality:invalid_product_codes", reasons)

    def test_dimensions_keys_and_foreign_keys(self):
        anonymous = self.base.copy()
        anonymous[6] = None
        self.write([self.base, self.base, anonymous])
        report = self.run_pipeline()
        self.assertEqual(report["reconciliation"]["counts"], {
            "staging_transaction": 3, "fact_transaction": 3, "dim_customer": 1,
            "dim_product": 1, "dim_date": 1, "dim_country": 1})
        self.assertEqual(report["duplicate_source_rows"], 1)
        with closing(sqlite3.connect(self.db)) as conn:
            conn.execute("PRAGMA foreign_keys=ON")
            self.assertFalse(conn.execute("PRAGMA foreign_key_check").fetchall())
            with self.assertRaises(sqlite3.IntegrityError):
                conn.execute("INSERT INTO dim_customer VALUES ('12345')")
            with self.assertRaises(sqlite3.IntegrityError):
                conn.execute("UPDATE fact_transaction SET product_key='nonexistent'")
            with self.assertRaises(sqlite3.IntegrityError):
                conn.execute("INSERT INTO fact_transaction SELECT * FROM fact_transaction LIMIT 1")

    def test_source_target_reconciliation(self):
        cancel = self.base.copy()
        cancel[0], cancel[3] = "C123456", -1
        zero = self.base.copy()
        zero[5] = 0
        invalid = self.base.copy()
        invalid[4] = "bad"
        self.write([self.base, cancel, zero, invalid])
        report = self.run_pipeline()
        self.assertEqual(report["raw_rows"], 4)
        self.assertEqual(report["valid_analytical_rows"], 2)
        self.assertEqual(report["excluded_rows"], 2)
        self.assertEqual(sum(report["primary_exclusion_reasons_disjoint"].values()), 2)
        self.assertEqual(report["analytical_totals"]["amount_micros"], 1235000)
        self.assertEqual(report["analytical_totals"]["invoices"], 2)
        self.assertEqual(report["reconciliation"]["status"], "passed")

    def test_idempotent_etl(self):
        self.write([self.base, self.base])
        first = self.run_pipeline()
        second = self.run_pipeline()
        self.assertEqual(first["reconciliation"], second["reconciliation"])
        self.assertEqual(first["source_sha256"], second["source_sha256"])

    def test_reconciliation_rejects_swapped_date_keys(self):
        other = self.base.copy()
        other[4] = datetime(2010, 1, 3)
        self.write([self.base, other])
        report = self.run_pipeline()
        with closing(sqlite3.connect(self.db)) as conn:
            conn.execute("""UPDATE fact_transaction SET date_key = CASE date_key
                WHEN 20100102 THEN 20100103 ELSE 20100102 END""")
            with self.assertRaisesRegex(ValueError, "timestamp/date dimension"):
                reconcile(conn, report)

    def test_failed_build_preserves_previous_database(self):
        self.write()
        self.run_pipeline()
        before = self.db.read_bytes()
        self.write(header=("Wrong",) + COLUMNS[1:])
        with self.assertRaises(ValueError):
            self.run_pipeline()
        self.assertEqual(self.db.read_bytes(), before)
        self.assertFalse(list(self.root.glob("retail_build_*")))

    def test_unrelated_database_protected(self):
        self.write()
        with closing(sqlite3.connect(self.db)) as conn:
            conn.execute("CREATE TABLE customers(id INTEGER)")
            conn.commit()
        with self.assertRaisesRegex(ValueError, "not owned"):
            self.run_pipeline()

    def test_empty_workbook(self):
        self.write([])
        report = self.run_pipeline()
        self.assertEqual(report["raw_rows"], 0)
        self.assertEqual(report["reconciliation"]["status"], "passed")


if __name__ == "__main__":
    unittest.main()
